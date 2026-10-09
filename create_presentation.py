import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Clean Green, White & Navy-Blue
    NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
    NAVY_MEDIUM = RGBColor(30, 41, 59)    # #1E293B
    GREEN_PRIMARY = RGBColor(21, 128, 61) # #15803D (Forest Green)
    GREEN_LIGHT = RGBColor(220, 252, 231) # #DCFCE7 (Soft mint tint)
    GREEN_ACCENT = RGBColor(34, 197, 94)  # #22C55E
    WHITE = RGBColor(255, 255, 255)       # #FFFFFF
    BG_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC (Clean background)
    SLATE_LIGHT = RGBColor(241, 245, 249) # #F1F5F9 (Card background)
    TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, tag_text, title_text, subtitle_text):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = GREEN_PRIMARY
        p.font.name = "Segoe UI"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        p_t.font.name = "Segoe UI"

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.45))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.size = Pt(14)
        p_s.font.color.rgb = TEXT_MUTED
        p_s.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 1: INTRODUCTION
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, BG_LIGHT)

    # Decorative top bar
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = GREEN_PRIMARY
    top_bar.line.color.rgb = GREEN_PRIMARY

    # Title Tag Badge
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(3.6), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = GREEN_LIGHT
    badge.line.color.rgb = GREEN_PRIMARY
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "SMART AGRICULTURE & CLEAN ENERGY"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = GREEN_PRIMARY
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.name = "Segoe UI"

    # Main Big Title
    title1_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.3))
    tf1 = title1_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Sun-Starved Tracker"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = NAVY_DARK
    p1.font.name = "Segoe UI"

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Intelligent Agri-Voltaics Optimizer"
    p1_sub.font.size = Pt(24)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = GREEN_PRIMARY
    p1_sub.font.name = "Segoe UI"

    # Main Idea Card
    main_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.2), Inches(11.7), Inches(2.2))
    main_card.fill.solid()
    main_card.fill.fore_color.rgb = WHITE
    main_card.line.color.rgb = BORDER_LIGHT
    main_card.line.width = Pt(1.5)

    tf_mc = main_card.text_frame
    tf_mc.word_wrap = True
    tf_mc.margin_left = Inches(0.4)
    tf_mc.margin_top = Inches(0.3)
    tf_mc.margin_right = Inches(0.4)

    p_mc_h = tf_mc.paragraphs[0]
    p_mc_h.text = "Core Concept: The Perfect Dual-Harvest Harmony"
    p_mc_h.font.size = Pt(18)
    p_mc_h.font.bold = True
    p_mc_h.font.color.rgb = NAVY_DARK
    p_mc_h.font.name = "Segoe UI"

    p_mc_b = tf_mc.add_paragraph()
    p_mc_b.space_before = Pt(8)
    p_mc_b.text = "Sun-Starved Tracker is an AI-powered IoT platform that balances solar power generation with healthy crop growth beneath solar panels. By dynamically tilting panels using real-time crop light needs and weather forecasts, the system guarantees zero crop starvation while extracting peak solar energy."
    p_mc_b.font.size = Pt(14)
    p_mc_b.font.color.rgb = NAVY_MEDIUM
    p_mc_b.font.name = "Segoe UI"

    # 4 Quick Pillar Badges at bottom
    pillars = [
        "🌱 Photosynthesis-Aware",
        "⚙️ Dual-Axis Kinematics",
        "🔋 Smart Battery Storage",
        "📱 Multi-Language Mobile Ready"
    ]
    card_w = Inches(2.7)
    gap = Inches(0.3)
    start_x = Inches(0.8)
    for i, pil in enumerate(pillars):
        x = start_x + i * (card_w + gap)
        pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(5.8), card_w, Inches(0.65))
        pill.fill.solid()
        pill.fill.fore_color.rgb = SLATE_LIGHT
        pill.line.color.rgb = BORDER_LIGHT
        pill.line.width = Pt(1)
        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = pil
        p_p.font.size = Pt(12)
        p_p.font.bold = True
        p_p.font.color.rgb = NAVY_DARK
        p_p.alignment = PP_ALIGN.CENTER
        p_p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 2: THE PROBLEM
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "THE CHALLENGE", "The Conflict: Crop Sunlight vs. Solar Energy",
               "Traditional static solar arrays create harsh tradeoffs on working farmlands.")

    col_width = Inches(5.7)
    col_height = Inches(4.3)
    top_y = Inches(2.2)

    # Column 1: Crop Sunlight Challenges (Green accent)
    card_crop = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_y, col_width, col_height)
    card_crop.fill.solid()
    card_crop.fill.fore_color.rgb = WHITE
    card_crop.line.color.rgb = GREEN_PRIMARY
    card_crop.line.width = Pt(2)

    tf_cc = card_crop.text_frame
    tf_cc.word_wrap = True
    tf_cc.margin_left = tf_cc.margin_right = Inches(0.35)
    tf_cc.margin_top = Inches(0.35)

    p_c1 = tf_cc.paragraphs[0]
    p_c1.text = "🌱 Crop Sunlight Needs"
    p_c1.font.size = Pt(20)
    p_c1.font.bold = True
    p_c1.font.color.rgb = GREEN_PRIMARY
    p_c1.font.name = "Segoe UI"

    crop_points = [
        ("Sun-Starvation:", "Fixed solar panels cast permanent dark shadows, depriving crops of essential PAR light."),
        ("Stunted Growth:", "Insufficient sunlight drops crop yields by 30% to 50% for high-light crops like tomatoes."),
        ("Microclimate Issues:", "Excess trapped dampness under unmoving panels leads to severe leaf mold and root rot."),
        ("Farmer Reluctance:", "Farmers resist solar installations fearing food security and revenue loss.")
    ]
    for title, desc in crop_points:
        p = tf_cc.add_paragraph()
        p.space_before = Pt(10)
        p.text = f"•  {title} {desc}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = NAVY_MEDIUM
        p.font.name = "Segoe UI"

    # Column 2: Solar Energy Challenges (Navy accent)
    card_solar = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), top_y, col_width, col_height)
    card_solar.fill.solid()
    card_solar.fill.fore_color.rgb = WHITE
    card_solar.line.color.rgb = NAVY_DARK
    card_solar.line.width = Pt(2)

    tf_cs = card_solar.text_frame
    tf_cs.word_wrap = True
    tf_cs.margin_left = tf_cs.margin_right = Inches(0.35)
    tf_cs.margin_top = Inches(0.35)

    p_s1 = tf_cs.paragraphs[0]
    p_s1.text = "☀️ Solar Energy Demands"
    p_s1.font.size = Pt(20)
    p_s1.font.bold = True
    p_s1.font.color.rgb = NAVY_DARK
    p_s1.font.name = "Segoe UI"

    solar_points = [
        ("Sub-Optimal Angles:", "Fixed-tilt panels only capture peak power during a narrow mid-day window (25% lost energy)."),
        ("Weather Vulnerability:", "Heavy hailstorms and severe wind gusts crack static glass and bend mounting brackets."),
        ("Land Competition:", "Large ground-mount solar plants lock up fertile agricultural soil for decades."),
        ("Grid Synchronization:", "Without smart local battery management, excess solar power is curtailed or wasted.")
    ]
    for title, desc in solar_points:
        p = tf_cs.add_paragraph()
        p.space_before = Pt(10)
        p.text = f"•  {title} {desc}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = NAVY_MEDIUM
        p.font.name = "Segoe UI"

    # Bottom takeaway bar
    bar2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.48))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = SLATE_LIGHT
    bar2.line.color.rgb = BORDER_LIGHT
    tf_b2 = bar2.text_frame
    p_b2 = tf_b2.paragraphs[0]
    p_b2.text = "💡 The Core Dilemma: Agriculture needs open sunlight. Solar needs full exposure. How can we have both?"
    p_b2.font.size = Pt(11)
    p_b2.font.bold = True
    p_b2.font.color.rgb = NAVY_DARK
    p_b2.alignment = PP_ALIGN.CENTER
    p_b2.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 3: THE SOLUTION
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "INNOVATIVE SOLUTION", "Sun-Starved Tracker: 4 Core Pillars",
               "An intelligent, closed-loop agrivoltaic management ecosystem built for farmers.")

    # 4 Modular Feature Cards (2x2 Grid)
    card_w3 = Inches(5.7)
    card_h3 = Inches(2.25)
    row1_y = Inches(2.1)
    row2_y = Inches(4.7)

    features = [
        {
            "col": Inches(0.8), "row": row1_y, "border": GREEN_PRIMARY,
            "title": "1. AI Crop Health & Light Optimization",
            "body": "Continuously measures PAR (Photosynthetically Active Radiation), canopy shade, and soil moisture. When crops are light-deficient, panels tilt away to allow filtered sunlight directly to root rows."
        },
        {
            "col": Inches(6.8), "row": row1_y, "border": NAVY_DARK,
            "title": "2. Dynamic Dual-Axis Kinematics",
            "body": "Smart motor actuators adjust panel elevation (0°–75°) and azimuth. Maximizes clean kilowatt-hours when crops are saturated, and provides protective shade during scorching heatwaves."
        },
        {
            "col": Inches(0.8), "row": row2_y, "border": NAVY_DARK,
            "title": "3. Smart Battery & Energy Management",
            "body": "Monitors dual LiFePO4 battery banks (SoC, cell health, charging cycles). Balances farm pump consumption, on-site storage, and profitable grid export without power interruptions."
        },
        {
            "col": Inches(6.8), "row": row2_y, "border": GREEN_PRIMARY,
            "title": "4. Automated Safety Stowing & Alerts",
            "body": "Autonomous storm response: stows flat at 0° during high winds to minimize drag (-86%), angles at 30° during rain for guided water harvesting, and sends instant SMS/voice alerts in 6 languages."
        }
    ]

    for f in features:
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, f["col"], f["row"], card_w3, card_h3)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = f["border"]
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p_h = tf.paragraphs[0]
        p_h.text = f["title"]
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = f["border"]
        p_h.font.name = "Segoe UI"

        p_b = tf.add_paragraph()
        p_b.space_before = Pt(6)
        p_b.text = f["body"]
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = NAVY_MEDIUM
        p_b.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 4: SYSTEM WORKING (4-STEP FLOW)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "HOW IT WORKS", "Simple 4-Step Closed-Loop Flow",
               "Automated sensing, AI intelligence, and precision mechanical actuation.")

    step_w = Inches(2.7)
    step_h = Inches(4.2)
    step_gap = Inches(0.3)
    start_x = Inches(0.8)
    step_y = Inches(2.2)

    steps = [
        {
            "num": "STEP 1", "icon": "📡", "name": "Sense",
            "color": GREEN_PRIMARY,
            "desc": "Ground & sky IoT sensors measure soil moisture, PAR sunlight, ambient temperature, battery SoC, and upcoming weather forecasts."
        },
        {
            "num": "STEP 2", "icon": "🧠", "name": "Analyze",
            "color": NAVY_DARK,
            "desc": "AI Optimizer calculates the real-time trade-off curve between crop photosynthesis requirements and peak solar power generation."
        },
        {
            "num": "STEP 3", "icon": "⚙️", "name": "Actuate",
            "color": GREEN_PRIMARY,
            "desc": "Quiet electric linear actuators tilt panels to the ideal elevation (e.g. 15° morning boost, 35° optimal track, or 0° storm stow)."
        },
        {
            "num": "STEP 4", "icon": "⚖️", "name": "Balance",
            "color": NAVY_DARK,
            "desc": "Crops receive optimal sunlight to thrive, panels generate peak clean energy, and stored battery power runs farm irrigation."
        }
    ]

    for i, s in enumerate(steps):
        x = start_x + i * (step_w + step_gap)
        scard = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, step_y, step_w, step_h)
        scard.fill.solid()
        scard.fill.fore_color.rgb = WHITE
        scard.line.color.rgb = s["color"]
        scard.line.width = Pt(2)

        tf = scard.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)

        p_num = tf.paragraphs[0]
        p_num.text = s["num"]
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = s["color"]
        p_num.font.name = "Segoe UI"

        p_name = tf.add_paragraph()
        p_name.space_before = Pt(4)
        p_name.text = f"{s['icon']} {s['name']}"
        p_name.font.size = Pt(20)
        p_name.font.bold = True
        p_name.font.color.rgb = NAVY_DARK
        p_name.font.name = "Segoe UI"

        p_desc = tf.add_paragraph()
        p_desc.space_before = Pt(14)
        p_desc.text = s["desc"]
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = NAVY_MEDIUM
        p_desc.font.name = "Segoe UI"

        # Connector indicator (except last step)
        if i < len(steps) - 1:
            arrow_x = x + step_w + Inches(0.05)
            arrow_box = slide4.shapes.add_textbox(arrow_x, step_y + Inches(1.8), Inches(0.25), Inches(0.4))
            tf_a = arrow_box.text_frame
            tf_a.margin_left = tf_a.margin_top = tf_a.margin_right = tf_a.margin_bottom = 0
            p_a = tf_a.paragraphs[0]
            p_a.text = "➔"
            p_a.font.size = Pt(18)
            p_a.font.bold = True
            p_a.font.color.rgb = GREEN_PRIMARY
            p_a.font.name = "Segoe UI"

    # Bottom workflow status note
    bar4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.5))
    bar4.fill.solid()
    bar4.fill.fore_color.rgb = GREEN_LIGHT
    bar4.line.color.rgb = GREEN_PRIMARY
    bar4.line.width = Pt(1)
    tf_b4 = bar4.text_frame
    p_b4 = tf_b4.paragraphs[0]
    p_b4.text = "⚡ Fully Autonomous Closed Loop: Re-evaluates every 15 minutes without requiring farmer manual intervention."
    p_b4.font.size = Pt(11)
    p_b4.font.bold = True
    p_b4.font.color.rgb = GREEN_PRIMARY
    p_b4.alignment = PP_ALIGN.CENTER
    p_b4.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 5: IMPACT & FUTURE VISION
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "IMPACT & VISION", "Projected Impact & Future Roadmap",
               "Delivering tangible financial and environmental results for farming communities.")

    # Top: 3 Big Metric Cards
    metric_w = Inches(3.7)
    metric_h = Inches(2.0)
    metric_gap = Inches(0.3)
    start_xm = Inches(0.8)
    top_ym = Inches(2.1)

    metrics = [
        {
            "val": "+35%", "label": "Higher Farmer Income",
            "desc": "Dual revenue stream: agricultural harvest sales plus surplus electricity grid feed-in tariffs.",
            "color": GREEN_PRIMARY
        },
        {
            "val": "+40%", "label": "Water Savings",
            "desc": "Controlled panel microclimate reduces soil moisture evaporation and scorching sun stress.",
            "color": NAVY_DARK
        },
        {
            "val": "100%", "label": "Severe Storm Protection",
            "desc": "Auto-stowing mechanism protects both expensive solar hardware and sensitive crop beds.",
            "color": GREEN_PRIMARY
        }
    ]

    for i, m in enumerate(metrics):
        x = start_xm + i * (metric_w + metric_gap)
        mcard = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top_ym, metric_w, metric_h)
        mcard.fill.solid()
        mcard.fill.fore_color.rgb = WHITE
        mcard.line.color.rgb = m["color"]
        mcard.line.width = Pt(1.5)

        tf = mcard.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)

        p_v = tf.paragraphs[0]
        p_v.text = m["val"]
        p_v.font.size = Pt(28)
        p_v.font.bold = True
        p_v.font.color.rgb = m["color"]
        p_v.font.name = "Segoe UI"

        p_l = tf.add_paragraph()
        p_l.text = m["label"]
        p_l.font.size = Pt(13)
        p_l.font.bold = True
        p_l.font.color.rgb = NAVY_DARK
        p_l.font.name = "Segoe UI"

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(4)
        p_d.text = m["desc"]
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.font.name = "Segoe UI"

    # Middle: Future Roadmap Box
    vision_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(1.75))
    vision_card.fill.solid()
    vision_card.fill.fore_color.rgb = WHITE
    vision_card.line.color.rgb = BORDER_LIGHT
    vision_card.line.width = Pt(1)

    tf_v = vision_card.text_frame
    tf_v.word_wrap = True
    tf_v.margin_left = tf_v.margin_right = Inches(0.3)
    tf_v.margin_top = Inches(0.2)

    p_vh = tf_v.paragraphs[0]
    p_vh.text = "🚀 Future Roadmap"
    p_vh.font.size = Pt(15)
    p_vh.font.bold = True
    p_vh.font.color.rgb = NAVY_DARK
    p_vh.font.name = "Segoe UI"

    roadmap_points = [
        "• Drone Multispectral Integration: Real-time NDVI scanning for automated pest and chlorophyll detection.",
        "• Village Microgrid Sharing: Peer-to-peer decentralized solar battery sharing among neighboring farmers.",
        "• AI Yield Forecasting: Machine learning models predicting crop yield based on seasonal agrivoltaic microclimates."
    ]
    for r in roadmap_points:
        p = tf_v.add_paragraph()
        p.space_before = Pt(3)
        p.text = r
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY_MEDIUM
        p.font.name = "Segoe UI"

    # Bottom: Thank You & Live Demo Callout
    thanks_bar = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.8))
    thanks_bar.fill.solid()
    thanks_bar.fill.fore_color.rgb = NAVY_DARK
    thanks_bar.line.color.rgb = NAVY_DARK

    tf_t = thanks_bar.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = Inches(0.3)
    tf_t.margin_top = Inches(0.15)

    p_th = tf_t.paragraphs[0]
    p_th.text = "Thank You! Questions & Live Demonstration"
    p_th.font.size = Pt(16)
    p_th.font.bold = True
    p_th.font.color.rgb = WHITE
    p_th.alignment = PP_ALIGN.CENTER
    p_th.font.name = "Segoe UI"

    p_sub_t = tf_t.add_paragraph()
    p_sub_t.text = "Live Web App: https://ayeshamundargi.github.io/sun-starved-tracker/  •  GitHub: github.com/ayeshamundargi/sun-starved-tracker"
    p_sub_t.font.size = Pt(10)
    p_sub_t.font.color.rgb = GREEN_LIGHT
    p_sub_t.alignment = PP_ALIGN.CENTER
    p_sub_t.font.name = "Segoe UI"

    output_path = os.path.abspath("Sun-Starved_Tracker_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    build_presentation()
